"""Native Excel PivotTables with cached results and worksheet-backed refresh.

The limited OOXML shape here (one row field, one column field, optional page
field, one count measure) is tested against desktop Excel. Counts are based
on purpose-built sources, never the CP Code x hostname All Data expansion.
"""
from collections import Counter
from openpyxl.pivot.cache import CacheDefinition, CacheSource, WorksheetSource, CacheField, SharedItems
from openpyxl.pivot.record import RecordList, Record
from openpyxl.pivot.fields import Text, Index
from openpyxl.pivot.table import (
    TableDefinition, PivotField, FieldItem, RowColField, RowColItem,
    DataField, PageField, Location, PivotTableStyle,
)
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def add_count_pivot(wb, source, title, row_field, col_field, measure, caption,
                    note, page_field=None):
    ws = wb.create_sheet(title)
    ws["A1"] = title
    ws["A1"].font = Font(size=18, bold=True, color="19304B")
    ws["A2"] = note
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 42
    ws.sheet_view.showGridLines = False
    headers = [str(c.value) for c in source[1]]
    rows = [[str(v) if v is not None else "" for v in row]
            for row in source.iter_rows(min_row=2, values_only=True)]
    ws.column_dimensions["A"].width = 35
    if not rows:
        ws["A4"] = "No source rows were collected for this view."
        return
    ri, ci, mi = map(headers.index, (row_field, col_field, measure))
    pi = headers.index(page_field) if page_field else None
    unique = [sorted({row[i] for row in rows}) for i in range(len(headers))]
    indices = [{v: i for i, v in enumerate(values)} for values in unique]
    cache = CacheDefinition(
        cacheSource=CacheSource(type="worksheet", worksheetSource=WorksheetSource(
            ref=source.dimensions, sheet=source.title)),
        cacheFields=[CacheField(name=header, sharedItems=SharedItems(
            _fields=[Text(v=v) for v in values], containsString=True,
            containsNonDate=True))
            for header, values in zip(headers, unique)],
        recordCount=len(rows), saveData=True, enableRefresh=True,
        refreshOnLoad=False, createdVersion=6, refreshedVersion=6, minRefreshableVersion=3)
    cache.records = RecordList(r=[Record(_fields=[Index(v=indices[i][value])
                              for i, value in enumerate(row)]) for row in rows])
    fields = []
    for i in range(len(headers)):
        axis = "axisRow" if i == ri else "axisCol" if i == ci else "axisPage" if i == pi else None
        fields.append(PivotField(axis=axis, dataField=i == mi, showAll=False,
            items=([FieldItem(x=j) for j in range(len(unique[i]))] + [FieldItem(t="default")]) if axis else []))
    start = 6 if page_field else 4
    end_col = get_column_letter(len(unique[ci]) + 2)
    end_row = start + len(unique[ri]) + 2
    pivot = TableDefinition(
        name="Audit" + "".join(c for c in title if c.isalnum()), cacheId=1 + sum(len(sheet._pivots) for sheet in wb.worksheets),
        dataCaption="Values", createdVersion=6, updatedVersion=6, minRefreshableVersion=3,
        location=Location(ref=f"A{start}:{end_col}{end_row}", firstHeaderRow=1,
            firstDataRow=2, firstDataCol=1, rowPageCount=1 if page_field else None,
            colPageCount=1 if page_field else None),
        pivotFields=fields,
        rowFields=[RowColField(x=ri)],
        rowItems=[RowColItem(x=[Index(v=i)]) for i in range(len(unique[ri]))] + [RowColItem(t="grand", x=[Index(v=0)])],
        colFields=[RowColField(x=ci)],
        colItems=[RowColItem(x=[Index(v=i)]) for i in range(len(unique[ci]))] + [RowColItem(t="grand", x=[Index(v=0)])],
        pageFields=[PageField(fld=pi, hier=-1)] if pi is not None else [],
        dataFields=[DataField(name=caption, fld=mi, subtotal="count", baseField=0, baseItem=0)],
        rowGrandTotals=True, colGrandTotals=True, compact=False, outline=True,
        outlineData=True, gridDropZones=False,
        pivotTableStyleInfo=PivotTableStyle(name="PivotStyleLight16", showRowHeaders=True,
            showColHeaders=True, showRowStripes=True, showColStripes=False, showLastColumn=True))
    pivot.cache = cache
    ws.add_pivot(pivot)
    # Display the initial result in every reader, including ones that do not
    # calculate PivotTables. Excel can refresh or rearrange using the cache.
    if page_field:
        ws["A4"], ws["B4"] = page_field, "(All)"
    ws.cell(start, 1, caption)
    ws.cell(start, 2, col_field)
    for j, value in enumerate([row_field] + unique[ci] + ["Grand Total"], 1):
        ws.cell(start + 1, j, value)
    counts = Counter((row[ri], row[ci]) for row in rows if row[mi] != "")
    for i, value in enumerate(unique[ri], start + 2):
        ws.cell(i, 1, value)
        for j, col in enumerate(unique[ci], 2):
            ws.cell(i, j, counts[value, col])
        ws.cell(i, len(unique[ci]) + 2, sum(counts[value, col] for col in unique[ci]))
    ws.cell(end_row, 1, "Grand Total")
    for j, col in enumerate(unique[ci], 2):
        ws.cell(end_row, j, sum(counts[value, col] for value in unique[ri]))
    ws.cell(end_row, len(unique[ci]) + 2, sum(counts.values()))
    for row in ws.iter_rows(min_row=start, max_row=end_row, max_col=len(unique[ci])+2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row in (start, start+1, end_row):
                cell.fill = PatternFill("solid", fgColor="E8F0FA")
                cell.font = Font(bold=True, color="19304B")
    for col in range(2, len(unique[ci])+3):
        ws.column_dimensions[get_column_letter(col)].width = 20
    ws.freeze_panes = f"B{start+2}"


def build_audit_pivots(wb, data):
    # Grouped findings have one row per finding/certificate/deployment group.
    findings = wb.create_sheet("Finding Pivot Source")
    findings.append(["Finding ID", "Severity", "Network", "Finding", "Rule references"])
    for index, item in enumerate(data.get("origin_action_groups", []), 1):
        findings.append([str(index), item.get("severity", "unknown"),
                        item.get("akamai_network", "Unknown"),
                        item.get("label") or item.get("finding", ""),
                        item.get("reference_count", 1)])
    findings.freeze_panes = "A2"
    findings.auto_filter.ref = findings.dimensions
    for row in findings:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    add_count_pivot(wb, wb["Edge TLS"], "Pivot - TLS", "Group", "Delivery mode",
                    "Property ID", "Configuration / network pairs",
                    "One property per network. Network filter starts at All (production + staging). Not active and Unknown remain visible. Use Excel PivotTable fields to change the view.",
                    page_field="Network")
    add_count_pivot(wb, wb["Edge Hostnames"], "Pivot - Certificates", "Certificate type", "Network",
                    "Hostname", "Hostname assignments",
                    "Counts property-hostname assignments per active network, not unique certificates or unique domains.")
    add_count_pivot(wb, findings, "Pivot - Origin Findings", "Severity", "Network",
                    "Finding ID", "Grouped findings",
                    "Counts grouped findings once per finding/certificate/deployment group. Repeated rule references are retained in Origin Actions and are not added to this count.")
