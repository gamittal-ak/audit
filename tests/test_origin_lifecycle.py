import asyncio
from copy import deepcopy
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook

from app.services.origin_findings import certificate_assessment, group_origin_actions, prepare_origin_report
from app.services.origin_cert_service import assess_origin, generate_recommendations
from app.services.excel_service import generate_excel, _HEADERS, _ACTION_HEADERS, _CERT_HEADERS

OBSERVED = datetime(2026,9,21,19,51,45,tzinfo=timezone.utc)


def certificate(days=24, lifetime=30, source="live_leaf", fingerprint="shared-certificate"):
    end=OBSERVED+timedelta(days=days,hours=2)
    return dict(source=source,collection_method="tls_handshake_no_verify",
        subject_cn="*.shield.example.com",issuer_org="Certainly",issuer_cn="Certainly Intermediate R1",
        sha256_fingerprint=fingerprint,not_before=(end-timedelta(days=lifetime)).isoformat(),
        not_after=end.isoformat(),days_remaining=days,observation_time=OBSERVED.isoformat(),
        is_expired=days<0,is_not_yet_valid=False,role="ca" if source=="configured_ca" else "leaf")


def origin(host="qa.shield.example.com",rule="default",cert=None):
    cert=cert or certificate()
    return dict(property_id="123",property_name="qa-video.example.com",property_version=303,
        akamai_network="PRODUCTION",rule_path=rule,origin_hostname=host,resolved_hostname=host,
        https_port=443,effective_sni=host,uses_https=True,origin_certs_to_honor="COMBO",
        configured_certificates=[],configured_cas=[],coverage_gaps=[],observation_status="observed",
        findings=[dict(finding="Certificate expiring within 30 days (live_leaf)",severity="warning",
                       days_remaining=24,cert_fingerprint=cert["sha256_fingerprint"],evidence="old")],
        live_certificates=[deepcopy(cert)])


def saved_report():
    origins=[origin(),origin("qa-vod.shield.example.com","default > VOD"),origin("qa-vod.shield.example.com","default > VOD-DR")]
    certs=[dict(o["live_certificates"][0],origin_hostname=o["resolved_hostname"],
                property_id=o["property_id"],property_name=o["property_name"],akamai_network=o["akamai_network"]) for o in origins]
    return dict(audit_timestamp=OBSERVED.isoformat(),report=[dict(groupname="Delivery",groupid="1",contractid="2",
        properties=[dict(id="123",name="qa-video.example.com",origin=["qa.shield.example.com"],cpcodes=[],hostnames=[])])],
        origin_inventory=origins,origin_certificates=certs,
        origin_actions=generate_recommendations(origins),
        origin_coverage=dict(total_origins=3,probed=3,total_certificates=3,expiring_30d=3,expired=0,actions_warning=3))


@pytest.mark.parametrize("days,expected",[(24,"info"),(11,"info"),(10,"warning"),(8,"warning"),(7,"critical"),(0,"critical"),(-1,"critical")])
def test_short_lived_boundaries(days,expected):
    result=certificate_assessment(certificate(days),origin())
    assert result["severity"]==expected
    assert result["renewal_status"]=="unverified"
    if days>10:
        assert result["code"]=="short_lived_review"
        assert "does not prove auto-renewal" in result["recommendation"]


def test_issuer_alone_and_missing_dates_do_not_downgrade():
    assert certificate_assessment(certificate(lifetime=90))["severity"]=="warning"
    cert=certificate();cert.pop("not_before")
    assert certificate_assessment(cert)["severity"]=="warning"
    cert["not_before"]="invalid date"
    assert certificate_assessment(cert)["severity"]=="warning"


@pytest.mark.parametrize("source,trust,pins",[
    ("configured_ca","COMBO",[]),
    ("configured_pin","SPECIFIC_CERTIFICATES",[]),
    ("live_leaf","SPECIFIC_CERTIFICATES",[]),
    ("live_leaf","COMBO",[{"sha256_fingerprint":"pin"}]),
])
def test_ca_and_pin_contexts_keep_warning(source,trust,pins):
    context=origin();context["origin_certs_to_honor"]=trust;context["configured_certificates"]=pins
    result=certificate_assessment(certificate(source=source),context)
    assert result["severity"]=="warning"
    if source!="configured_ca":
        assert "pinning" in result["recommendation"]


def test_not_yet_valid_is_critical_and_exact_expiry_is_expired():
    cert=certificate();cert["not_before"]=(OBSERVED+timedelta(hours=1)).isoformat()
    assert certificate_assessment(cert)["code"]=="not_yet_valid"
    cert["not_before"]=(OBSERVED-timedelta(days=30)).isoformat()
    cert["not_after"]=OBSERVED.isoformat()
    result=certificate_assessment(cert)
    assert result["code"]=="expired" and result["severity"]=="critical"


def test_snapshot_reinterpretation_preserves_evidence_and_groups_references():
    raw=saved_report();before=deepcopy(raw)
    updated=prepare_origin_report(raw)
    assert raw==before
    assert len(updated["origin_actions"])==3  # Excel/raw action references retained.
    assert len(updated["origin_action_groups"])==1
    finding=updated["origin_action_groups"][0]
    assert finding["severity"]=="info" and finding["days_remaining"]==24
    assert finding["reference_count"]==3
    assert set(finding["origin_hostnames"])=={"qa.shield.example.com","qa-vod.shield.example.com"}
    assert {r["rule_path"] for r in finding["references"]}=={"default","default > VOD","default > VOD-DR"}
    assert updated["origin_findings_summary"]["renewal_review"]==1
    assert updated["origin_findings_summary"]["unique_certificates"]==1
    assert len(updated["origin_certificates"])==3
    assert prepare_origin_report(updated)==updated


def test_no_certificate_grouping_without_identity_or_across_distinct_fingerprints():
    actions=prepare_origin_report(saved_report())["origin_actions"]
    actions[1]["cert_fingerprint"]="another"
    assert len(group_origin_actions(actions))==2
    for action in actions:
        action.pop("cert_fingerprint")
    assert len(group_origin_actions(actions))==3


def test_network_and_trust_differences_are_preserved():
    actions=prepare_origin_report(saved_report())["origin_actions"]
    actions[1]["akamai_network"]="STAGING"
    actions[2]["recommendation"]="Different trust instructions"
    assert len(group_origin_actions(actions))==3


def test_repeated_endpoint_observations_do_not_duplicate_findings():
    item=origin();cert=certificate()
    probe={"status":"ok","certificates":[cert,deepcopy(cert)],"probes":[{"certificates":[cert]},{"certificates":[cert]}]}
    assessed=assess_origin(item,probe)
    assert len(assessed["live_certificates"])==2  # Raw evidence remains.
    assert len(assessed["findings"])==1
    assert generate_recommendations([assessed])[0]["severity"]=="info"


def test_mixed_endpoints_warning_survives_saved_reinterpretation():
    raw=saved_report()
    raw["origin_inventory"][0]["findings"].append(dict(severity="warning",finding="Different certificates observed across endpoints",evidence="2 distinct fingerprints"))
    result=prepare_origin_report(raw)
    assert any(a["finding"]=="Different certificates observed across endpoints" and a["severity"]=="warning" for a in result["origin_actions"])


def test_missing_legacy_origin_data_is_safe():
    result=prepare_origin_report({"report":[]})
    assert result["origin_coverage"]=={}
    assert result["origin_actions"]==[]


def test_excel_retains_sheets_columns_rows_and_corrected_recommendations(tmp_path):
    raw=saved_report();path=tmp_path/"report.json";out=tmp_path/"report.xlsx"
    path.write_text(json.dumps(raw));before=path.read_bytes()
    generate_excel(str(path),str(out))
    wb=load_workbook(out)
    assert wb.sheetnames[:5]==["Summary","All Data","Origins","Origin Certificates","Origin Actions"]
    assert "Edge TLS" in wb.sheetnames and "Edge Hostnames" in wb.sheetnames
    assert [c.value for c in wb["All Data"][1]]==_HEADERS
    assert [c.value for c in wb["Origin Actions"][1]]==_ACTION_HEADERS
    assert [c.value for c in wb["Origin Certificates"][1]]==_CERT_HEADERS
    rows=list(wb["Origin Actions"].iter_rows(min_row=2,values_only=True))
    assert len(rows)==3
    assert all(row[0]=="info" and "renewal status unverified" in row[7] for row in rows)
    assert all("automated renewal" in row[10] for row in rows)
    assert wb["Origin Certificates"]["A2"].fill.fgColor.rgb not in ("00FF8C00","00CC0000")
    summary={row[0].value:row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Renewal reviews (rule references)"]==3
    assert path.read_bytes()==before
    wb.close()


def test_saved_download_uses_temporary_workbook_and_cleans_up(monkeypatch,tmp_path):
    from app.routers import downloads
    raw=saved_report();path=tmp_path/"report.json";original=tmp_path/"report.xlsx"
    path.write_text(json.dumps(raw));original.write_bytes(b"original")
    monkeypatch.setattr(downloads.celery_app,"AsyncResult",lambda _:SimpleNamespace(state="SUCCESS",result={"json_path":str(path),"xlsx_path":str(original)}))
    response=asyncio.run(downloads.download_xlsx("test",SimpleNamespace(session={"authenticated":True})))
    temporary=Path(response.path)
    assert temporary!=original and temporary.exists()
    assert original.read_bytes()==b"original"
    assert path.read_text()==json.dumps(raw)
    asyncio.run(response.background())
    assert not temporary.exists()


def test_download_falls_back_to_saved_metadata(monkeypatch,tmp_path):
    from app.routers import downloads
    path=tmp_path/"legacy.xlsx";path.write_bytes(b"original")
    monkeypatch.setattr(downloads.celery_app,"AsyncResult",lambda _:SimpleNamespace(state="PENDING"))
    connection=SimpleNamespace(hgetall=AsyncMock(return_value={"xlsx_path":str(path)}),aclose=AsyncMock())
    monkeypatch.setattr(downloads.aioredis,"from_url",AsyncMock(return_value=connection))
    response=asyncio.run(downloads.download_xlsx("test",SimpleNamespace(session={"authenticated":True})))
    assert Path(response.path)==path

def test_identical_connection_failures_group_by_endpoint_with_all_rules():
    data=saved_report()
    for item in data["origin_inventory"]:
        item["live_certificates"]=[]
        item["findings"]=[]
        item["observation_status"]="timeout"
        item["observation_error"]="Connection timed out"
    result=prepare_origin_report(data)
    assert len(result["origin_action_groups"])==2
    assert sorted(g["reference_count"] for g in result["origin_action_groups"])==[1,2]

def test_legacy_partial_inventory_does_not_drop_unassessed_actions():
    data=saved_report()
    extra=origin("legacy.example.com","default > legacy")
    for key in ("live_certificates","configured_certificates","configured_cas","findings"):
        extra.pop(key,None)
    data["origin_inventory"].append(extra)
    data["origin_actions"].append(dict(property_id="123",property_version=303,property_name="qa-video.example.com",
        akamai_network="PRODUCTION",rule_path="default > legacy",origin_hostname="legacy.example.com",
        finding="Legacy warning",severity="warning",recommendation="Investigate",evidence="Saved evidence"))
    result=prepare_origin_report(data)
    assert any(a["finding"]=="Legacy warning" for a in result["origin_actions"])
