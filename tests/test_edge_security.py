import asyncio
import pytest
import copy
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock
from openpyxl import load_workbook

from app.services.edge_security import classify_hostname, collect_property_security, prepare_edge_report, summarize
from app.services.akamai_client import AkamaiClient
from app.services.excel_service import generate_excel


def row(**kwargs):
    return dict(cnameFrom="www.example.com",cnameTo="www.example.com.edgekey.net",
                edgeHostnameId="ehn_42",certProvisioningType="CPS_MANAGED",**kwargs)


def test_suffix_is_not_tls_or_shared_evidence():
    for suffix in ("edgekey.net","edgesuite.net","akamaized.net"):
        h=row();h["cnameTo"]="www.example.com."+suffix
        out=classify_hostname(h,{},"PRODUCTION")
        assert out["tls_mode"]=="Unknown"
        assert out["certificate_type"]=="CPS-managed"
        assert out["protocol"]=="Unknown"


def test_legacy_standard_tls_on_edgekey_and_secure_false_never_imply_http_only():
    h=row()
    out=classify_hostname(h,{"securityType":"STANDARD-TLS","secure":False},"PRODUCTION")
    assert out["tls_mode"]=="sTLS" and out["protocol"]=="Unknown"
    h["certStatus"]={"production":[{"status":"STALLED"}]}
    assert classify_hostname(h,{},"PRODUCTION")["protocol"]=="Unknown"


def akamaized(name, cname_type, provision="CPS_MANAGED"):
    item = {"cnameFrom": name, "cnameTo": name, "edgeHostnameId": "1", "certProvisioningType": provision}
    if cname_type:
        item["cnameType"] = cname_type
    return classify_hostname(item, {"securityType": "STANDARD-TLS"}, "PRODUCTION")


@pytest.mark.parametrize("cname_type,cert,mode,protocol", [
    ("SHARED_CERT", "Akamai shared", "sTLS", "Shared HTTPS configured"),
    ("CUSTOM", "No certificate", "HTTP-only", "HTTP-only (Property Manager: no certificate)"),
    ("EDGE_HOSTNAME", "No certificate", "HTTP-only", "HTTP-only (Property Manager: no certificate)"),
    (None, "Unknown", "sTLS", "Unknown"),
])
def test_akamaized_certificate_follows_papi_cname_type(cname_type, cert, mode, protocol):
    row = akamaized("sports.akamaized.net", cname_type)
    assert row["certificate_type"] == cert
    assert row.get("delivery_mode", row["tls_mode"]) == mode
    assert row["protocol"] == protocol


def test_edge_hostname_cname_type_alone_is_not_http_only():
    for provision in ("DEFAULT", "CPS_MANAGED"):
        item = {"cnameFrom": "www.example.com", "cnameTo": "www.example.com.edgekey.net",
                "cnameType": "EDGE_HOSTNAME", "certProvisioningType": provision}
        row = classify_hostname(item, {"securityType": "ENHANCED-TLS"}, "PRODUCTION")
        assert "delivery_mode" not in row and row["certificate_type"] != "No certificate"
    assert akamaized("a.video.akamaized.net", "CUSTOM").get("delivery_mode") is None
    assert akamaized("sports.akamaized.net", "CUSTOM", "DEFAULT")["certificate_type"] == "Default DV"


def test_saved_name_based_shared_rows_are_reclassified_on_display():
    old = akamaized("sports.akamaized.net", "SHARED_CERT")
    old["raw_hostname"]["cnameType"] = "CUSTOM"
    kept = akamaized("weather.akamaized.net", "SHARED_CERT")
    data = {"report": [{"properties": [{"productionVersion": 3, "stagingVersion": None,
            "edge_security": {"PRODUCTION": dict(**summarize([old, kept]), version=3)}}]}]}
    out = prepare_edge_report(data)
    rows = out["report"][0]["properties"][0]["edge_security"]["PRODUCTION"]
    assert [r["certificate_type"] for r in rows["hostnames"]] == ["No certificate", "Akamai shared"]
    assert rows["label"] == "Mixed" and rows["version"] == 3
    assert out["edge_security_summary"]["PRODUCTION"]["HTTP-only"] == 1
    assert data["report"][0]["properties"][0]["edge_security"]["PRODUCTION"]["hostnames"][0]["certificate_type"] == "Akamai shared"


def test_hostname_pattern_alone_never_proves_shared_or_http_only():
    for name in ("video.akamaized.net","a.video.akamaized.net","video.akamaized.net.attacker.example"):
        h=row();h["cnameFrom"]=name;h["cnameTo"]=name
        out=classify_hostname(h,{},"PRODUCTION")
        assert out["certificate_type"]!="Akamai shared" and "delivery_mode" not in out
    for name in ("a.video.akamaized.net","video.akamaized.net.attacker.example"):
        h=row(cnameType="CUSTOM");h["cnameFrom"]=name;h["cnameTo"]=name
        assert "delivery_mode" not in classify_hostname(h,{},"PRODUCTION")
    h=row(cnameType="CUSTOM");h["cnameFrom"]="VIDEO.AKAMAIZED.NET.";h["cnameTo"]="video.akamaized.net"
    assert classify_hostname(h,{},"PRODUCTION")["delivery_mode"]=="HTTP-only"


def test_network_specific_deployment_status_and_default_pending():
    h=row(certStatus={"production":[{"status":"DEPLOYED"}],"staging":[{"status":"PENDING"}]})
    h["certProvisioningType"]="DEFAULT"
    assert classify_hostname(h,{},"PRODUCTION")["protocol"]=="HTTPS certificate deployed"
    assert "pending" in classify_hostname(h,{},"STAGING")["protocol"]


def test_same_version_reuses_hostnames_without_losing_both_networks():
    client=SimpleNamespace(get_edge_security_catalog=AsyncMock(return_value={"edgeHostnames":[
        {"edgeHostnameId":42,"securityType":"ENHANCED-TLS"}]}), get_hostnames=AsyncMock())
    source={"hostnames":{"items":[row(certStatus={"production":[{"status":"DEPLOYED"}]})]}}
    result=asyncio.run(collect_property_security(client,"acct","c","g","p",
        {"productionVersion":4,"stagingVersion":4},"4",source))
    assert set(result)=={"PRODUCTION","STAGING"}
    assert result["STAGING"]["label"]=="eTLS"
    assert result["PRODUCTION"]["hostnames"][0]["protocol"]=="HTTPS certificate deployed"
    assert result["STAGING"]["hostnames"][0]["protocol"]=="Unknown"
    client.get_hostnames.assert_not_called()


def test_failed_stage_and_hapi_do_not_drop_property_or_guess():
    client=SimpleNamespace(get_edge_security_catalog=AsyncMock(side_effect=PermissionError()),
                           get_hostnames=AsyncMock(side_effect=PermissionError()))
    result=asyncio.run(collect_property_security(client,"a","c","g","p",
        {"productionVersion":4,"stagingVersion":5},"4",{"hostnames":{"items":[row()]}}))
    assert result["PRODUCTION"]["label"]=="Unknown"
    assert result["PRODUCTION"]["hostname_count"]==1
    assert result["STAGING"]["status"]=="unavailable"


def test_bucket_separates_active_mappings():
    client=SimpleNamespace(get_edge_security_catalog=AsyncMock(return_value={"edgeHostnames":[]}),
        get_active_hostnames=AsyncMock(return_value=[{"cnameFrom":"example.com",
            "productionCnameTo":"prod.edgekey.net","productionCertType":"DEFAULT",
            "stagingCnameTo":"stage.edgesuite.net","stagingCertType":"CPS_MANAGED"}]))
    result=asyncio.run(collect_property_security(client,"a","c","g","p",
        {"useHostnameBucket":True,"productionVersion":4,"stagingVersion":4},"4",{}))
    assert result["PRODUCTION"]["hostnames"][0]["edge_hostname"]=="prod.edgekey.net"
    assert result["STAGING"]["hostnames"][0]["certificate_type"]=="CPS-managed"


def test_unknown_partial_mixed_counts_and_legacy_do_not_mutate():
    first=classify_hostname(row(),{"securityType":"ENHANCED-TLS"},"PRODUCTION")
    second=classify_hostname(row(),{"securityType":"STANDARD-TLS"},"PRODUCTION")
    unknown=classify_hostname(row(),{},"PRODUCTION")
    item=summarize([first,second,unknown])
    assert item["label"]=="Mixed + Unknown"
    raw={"report":[{"properties":[{"id":"p","productionVersion":4,"stagingVersion":None,
          "cert_type":"Shared Akamai Cert","edge_security":{"PRODUCTION":item}}]}]}
    before=copy.deepcopy(raw)
    new=prepare_edge_report(raw)
    assert raw==before
    counts=new["edge_security_summary"]["PRODUCTION"]
    assert counts["sTLS"]==counts["eTLS"]==counts["Mixed"]==counts["Unknown"]==1
    assert new["report"][0]["properties"][0]["edge_security"]["STAGING"]["status"]=="inactive"
    assert prepare_edge_report(new)==new
    raw["report"][0]["properties"][0].pop("edge_security")
    legacy=prepare_edge_report(raw)["report"][0]["properties"][0]
    assert legacy["edge_security"]["PRODUCTION"]["status"]=="not_collected"
    assert legacy["edge_certificate_label"]=="Unknown"
    assert legacy["cert_type"]=="Shared Akamai Cert"


def test_hapi_catalog_single_flight_including_failure():
    async def run():
        client=AkamaiClient("https://example.com",None,asyncio.Semaphore(2))
        client._get=AsyncMock(return_value={"edgeHostnames":[]})
        await asyncio.gather(*(client.get_edge_security_catalog("a") for _ in range(5)))
        assert client._get.await_count==1
        await client.get_edge_security_catalog("b")
        assert client._get.await_count==2
    asyncio.run(run())


def test_excel_has_network_rows_and_keeps_saved_evidence(tmp_path):
    source=tmp_path/"report.json";target=tmp_path/"report.xlsx"
    raw={"report":[{"groupname":"Example","groupid":"g","contractid":"c","properties":[
        {"id":"p","name":"Example","productionVersion":1,"stagingVersion":2,
         "hostnames":[],"cpcodes":[],"cert_type":"Shared Akamai Cert"}]}]}
    source.write_text(json.dumps(raw));before=source.read_bytes()
    generate_excel(str(source),str(target))
    wb=load_workbook(target)
    rows=list(wb["Edge TLS"].values)
    assert len(rows)==3
    assert {row[4] for row in rows[1:]}=={"PRODUCTION","STAGING"}
    assert all(row[6]=="Unknown" for row in rows[1:])
    assert source.read_bytes()==before
    assert wb["All Data"].max_column==47
