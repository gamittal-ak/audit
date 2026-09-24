import json
from collections import Counter
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook, Workbook
from app.services.excel_service import generate_excel
from app.services.pivot_service import add_count_pivot
from test_report_ui import report_context


def test_pivot_caches_sources_and_totals_are_distinct(tmp_path):
    ctx=report_context()
    source=tmp_path/"report.json";out=tmp_path/"report.xlsx"
    source.write_text(json.dumps(ctx))
    generate_excel(str(source),str(out))
    book=load_workbook(out)
    pivots=[book[name]._pivots[0] for name in ("Pivot - TLS","Pivot - Certificates","Pivot - Origin Findings")]
    assert len({p.cacheId for p in pivots})==3
    assert [p.cache.cacheSource.worksheetSource.sheet for p in pivots]==["Edge TLS","Edge Hostnames","Finding Pivot Source"]
    assert [p.cache.recordCount for p in pivots[:2]]==[6,6]
    assert len(pivots[0].cache.records.r)==6
    for name in ("Pivot - TLS","Pivot - Certificates"):
        ws=book[name];corner=ws._pivots[0].location.ref.split(":")[1]
        assert ws[corner].value==6
    assert book["Pivot - TLS"]._pivots[0].pageFields[0].fld==4
    with ZipFile(out) as archive:
        ns={"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        for name in archive.namelist():
            if "pivotCacheDefinition" in name and name.endswith(".xml"):
                tree=ET.fromstring(archive.read(name))
                assert all(node.get("containsSemiMixedTypes")!="0" for node in tree.findall(".//s:sharedItems",ns))


def test_empty_pivot_source_has_no_phantom_records():
    book=Workbook();source=book.active;source.append(["Row","Column","ID"])
    add_count_pivot(book,source,"Empty","Row","Column","ID","Count","No data")
    assert not book["Empty"]._pivots
    assert "No source rows" in book["Empty"]["A4"].value


def test_external_labels_are_text_not_formulas(tmp_path):
    book=Workbook();source=book.active;source.title="Source"
    source.append(["Row","Column","ID"]);source.append(["=HYPERLINK()","@test","id"])
    add_count_pivot(book,source,"Pivot","Row","Column","ID","Count","One row")
    assert book["Pivot"]["A6"].data_type=="s"
    path=tmp_path/"pivot.xlsx";book.save(path)
    assert load_workbook(path)["Pivot"]["A6"].value=="=HYPERLINK()"
