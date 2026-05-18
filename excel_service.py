"""
Converts a report JSON dict to an Excel .xlsx file.
Ported from json_to_excel.py — duplicate format_worksheet removed,
hostname handling fixed (uses dict .get() correctly).
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

# ------------------------------------------------------------------ styles
_header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
_header_font = Font(color="FFFFFF", bold=True)
_alt_fill = PatternFill(start_color="E6F0F8", end_color="E6F0F8", fill_type="solid")
_group_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
_prop_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
_warn_fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

_HEADERS = [
    "Group Name", "Group ID", "Parent Group ID", "Contract ID",
    "Property ID", "Property Name", "Origin", "Latest Version",
    "Staging Version", "Production Version", "CP Code", "Description",
    "Product", "Offload %", "Edge GB", "Midgress GB", "Origin GB",
    "Cache Hit %", "CNAME From", "CNAME To", "Cert Expiry",
    "Cert Issuer", "Min TLS Version", "Origin Characteristics", "SRO",
    "Site Shield", "Advanced Override", "Custom Override",
    "Custom Behavior Count", "CW/QR", "Total Rules", "Max Depth",
    "Behavior Count", "Last Activated", "Activated By",
]


def generate_excel(json_path: str, xlsx_path: str) -> str:
    """
    Read the report JSON at json_path, produce an Excel workbook at xlsx_path.
    Returns xlsx_path.
    """
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    wb = Workbook()
    ws = wb.active
    ws.title = "All Data"
    ws.append(_HEADERS)

    for group in data.get("report", []):
        group_name = group.get("groupname", "")
        group_id = group.get("groupid", "")
        parent_group_id = group.get("parentgroupid", "")
        contract_id = group.get("contractid", "")
        properties = group.get("properties", [])

        if not properties:
            ws.append([group_name, group_id, parent_group_id, contract_id] + [""] * (len(_HEADERS) - 4))
            continue

        for prop in properties:
            base = _prop_base(group_name, group_id, parent_group_id, contract_id, prop)
            cpcodes = prop.get("cpcodes", [])
            hostnames = prop.get("hostnames", [])

            flags = _prop_flags(prop)

            cert_tls = _cert_tls_cols(prop)

            if not cpcodes and not hostnames:
                ws.append(base + [""] * 8 + ["", ""] + cert_tls + flags)
                continue

            if not hostnames:
                for entry in _expand_cpcodes(cpcodes):
                    ws.append(base + entry + ["", ""] + cert_tls + flags)
                continue

            if not cpcodes:
                for h in hostnames:
                    ws.append(base + [""] * 8 + _hostname_cols(h) + cert_tls + flags)
                continue

            for entry in _expand_cpcodes(cpcodes):
                for h in hostnames:
                    ws.append(base + entry + _hostname_cols(h) + cert_tls + flags)

    _format_worksheet(ws)
    wb.save(xlsx_path)
    return xlsx_path


# ------------------------------------------------------------------ helpers

def _prop_base(group_name, group_id, parent_group_id, contract_id, prop) -> list:
    """Columns 1-10: group/property identity fields."""
    origin = prop.get("origin") or []
    if isinstance(origin, list):
        origin_str = ", ".join(str(x) for x in origin)
    else:
        origin_str = str(origin)

    return [
        group_name, group_id, parent_group_id, contract_id,
        prop.get("id", ""),
        prop.get("name", ""),
        origin_str,
        prop.get("latestVersion", ""),
        prop.get("stagingVersion", ""),
        prop.get("productionVersion", ""),
    ]


def _prop_flags(prop) -> list:
    """Analysis flag fields, appended after cpcode/traffic/hostname/cert columns."""
    cw_qr = prop.get("CW_QR") or []
    cw_qr_str = ", ".join(str(x) for x in cw_qr) if isinstance(cw_qr, list) else str(cw_qr)
    return [
        "Yes" if prop.get("originCharacteristics") else "No",
        prop.get("sro", ""),
        prop.get("site_shield", ""),
        prop.get("adv_override_exists", ""),
        prop.get("custom_override_exists", ""),
        prop.get("count_custom_behavior", ""),
        cw_qr_str,
        prop.get("total_rules", ""),
        prop.get("max_depth", ""),
        prop.get("behavior_count", ""),
        prop.get("last_activated", ""),
        prop.get("activated_by", ""),
    ]


def _expand_cpcodes(cpcodes: list) -> list:
    """Expand each cpcode entry into a [cpcode, description, product, traffic...] row fragment."""
    rows = []
    for entry in cpcodes:
        ids = entry.get("cpcode") or []
        descs = entry.get("description") or []
        prods = entry.get("product") or []
        if not ids:
            rows.append([""] * 8)
            continue
        traffic = entry.get("traffic") or {}
        count = max(len(ids), 1)
        for i in range(count):
            rows.append([
                ids[i] if i < len(ids) else "",
                descs[i] if i < len(descs) else "",
                prods[i] if i < len(prods) else "",
                traffic.get("bytesOffload", ""),
                traffic.get("edgeBytes", ""),
                traffic.get("midgressBytes", ""),
                traffic.get("originBytes", ""),
                traffic.get("cacheHitPct", ""),
            ])
    return rows or [[""] * 8]


def _hostname_cols(h) -> list:
    if isinstance(h, dict):
        return [h.get("cnameFrom", ""), h.get("cnameTo", "")]
    return [str(h), ""]


def _cert_tls_cols(prop) -> list:
    """Columns for Cert Expiry, Cert Issuer, Min TLS Version."""
    return [
        prop.get("cert_expiry", ""),
        prop.get("cert_issuer", ""),
        prop.get("min_tls", ""),
    ]


def _format_worksheet(ws):
    # Header row styling
    for cell in ws[1]:
        cell.fill = _header_fill
        cell.font = _header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _border

    ws.freeze_panes = "A2"

    # Auto-width (capped at 40)
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=0)
        ws.column_dimensions[col[0].column_letter].width = min((max_len + 2) * 1.2, 40)

    # Determine Edge GB column index dynamically
    edge_gb_idx = _HEADERS.index("Edge GB")

    # Row formatting
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), 2):
        group_name = row[0].value
        property_id = row[4].value
        # Check for zero traffic — flag red if Edge GB is 0 or empty
        edge_val = row[edge_gb_idx].value
        is_zero_traffic = property_id and (edge_val is None or edge_val == "" or edge_val == 0 or edge_val == 0.0)
        if is_zero_traffic:
            row_fill = _warn_fill
        else:
            row_fill = _group_fill if not property_id else _prop_fill
        for cell in row:
            cell.fill = row_fill
            cell.border = _border
            cell.alignment = Alignment(vertical="center")

    # Excel table
    tab = Table(displayName="ReportData", ref=ws.dimensions)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tab)
